#!/usr/bin/env python3
"""Minimal Markdown -> .docx converter (python-docx), tuned for the v2.1 document.

Handles: # headings (1-4), paragraphs, bullet and numbered lists (one level),
pipe tables (with header row), blockquotes, fenced code blocks (monospace,
preserved whitespace), horizontal rules, and inline **bold**, *italic*, `code`.
"""
import re
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*(?!\s).+?(?<!\s)\*|`[^`]+`)")


def add_inline(paragraph, text, base_bold=False, base_italic=False, size=None, font=None):
    """Add text with inline markdown formatting to a paragraph."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            _run(paragraph, text[pos:m.start()], base_bold, base_italic, size, font)
        tok = m.group(0)
        if tok.startswith("**"):
            _run(paragraph, tok[2:-2], True, base_italic, size, font)
        elif tok.startswith("`"):
            _run(paragraph, tok[1:-1], base_bold, base_italic, size, "Consolas")
        else:
            _run(paragraph, tok[1:-1], base_bold, True, size, font)
        pos = m.end()
    if pos < len(text):
        _run(paragraph, text[pos:], base_bold, base_italic, size, font)


def _run(paragraph, text, bold, italic, size, font):
    if not text:
        return
    r = paragraph.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if font:
        r.font.name = font
        rpr = r._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), font)
        rfonts.set(qn("w:cs"), font)


def set_cell_shading(cell, hex_fill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def set_table_borders(table):
    tbl = table._element
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "999999")
        borders.append(el)
    tblPr.append(borders)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    # split on unescaped pipes
    cells = re.split(r"(?<!\\)\|", line)
    return [c.strip().replace("\\|", "|") for c in cells]


def is_sep_row(line):
    s = line.strip()
    return bool(re.fullmatch(r"\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?", s))


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # Page setup and base style
    for section in doc.sections:
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for lvl, sz in ((1, 20), (2, 15), (3, 12.5), (4, 11)):
        st = doc.styles[f"Heading {lvl}"]
        st.font.name = "Calibri"
        st.font.size = Pt(sz)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        st.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
        st.paragraph_format.space_after = Pt(6)

    i = 0
    n = len(lines)
    para_buf = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            text = " ".join(s.strip() for s in para_buf)
            p = doc.add_paragraph()
            add_inline(p, text)
            para_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # fenced code block
        if stripped.startswith("```"):
            flush_para()
            i += 1
            code_lines = []
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i])
                i += 1
            i += 1  # closing fence
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.0
            for k, cl in enumerate(code_lines):
                r = p.add_run(cl.rstrip("\n"))
                r.font.name = "Consolas"
                r.font.size = Pt(7.5)
                rpr = r._element.get_or_add_rPr()
                rfonts = rpr.find(qn("w:rFonts"))
                if rfonts is None:
                    rfonts = OxmlElement("w:rFonts")
                    rpr.append(rfonts)
                rfonts.set(qn("w:eastAsia"), "Consolas")
                if k < len(code_lines) - 1:
                    r.add_break()
            continue

        # blank line
        if not stripped:
            flush_para()
            i += 1
            continue

        # horizontal rule
        if re.fullmatch(r"-{3,}|\*{3,}|_{3,}", stripped):
            flush_para()
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            pPr = p._element.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "999999")
            pbdr.append(bottom)
            pPr.append(pbdr)
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_para()
            level = min(len(m.group(1)), 4)
            h = doc.add_heading(level=level)
            add_inline(h, m.group(2).strip())
            i += 1
            continue

        # table
        if stripped.startswith("|") and i + 1 < n and is_sep_row(lines[i + 1]):
            flush_para()
            header = split_row(stripped)
            rows = []
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i].strip()))
                i += 1
            ncols = len(header)
            table = doc.add_table(rows=1 + len(rows), cols=ncols)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            set_table_borders(table)
            for c, txt in enumerate(header):
                cell = table.rows[0].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                add_inline(p, txt, base_bold=True, size=9)
                set_cell_shading(cell, "E8ECF2")
            for r_idx, row in enumerate(rows, start=1):
                row = (row + [""] * ncols)[:ncols]
                for c, txt in enumerate(row):
                    cell = table.rows[r_idx].cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    add_inline(p, txt, size=9)
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # blockquote
        if stripped.startswith(">"):
            flush_para()
            quote_lines = []
            while i < n and lines[i].strip().startswith(">"):
                quote_lines.append(lines[i].strip()[1:].strip())
                i += 1
            text = " ".join(q for q in quote_lines if q)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(1.0)
            p.paragraph_format.right_indent = Cm(0.6)
            pPr = p._element.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "12")
            left.set(qn("w:space"), "8")
            left.set(qn("w:color"), "1F2A44")
            pbdr.append(left)
            pPr.append(pbdr)
            add_inline(p, text)
            continue

        # bullet list
        m = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m:
            flush_para()
            while i < n:
                mm = re.match(r"^\s*[-*+]\s+(.*)$", lines[i])
                if not mm:
                    break
                item = mm.group(1)
                i += 1
                # continuation lines (indented, non-blank, not a new bullet)
                while i < n and lines[i].strip() and lines[i].startswith("  ") and not re.match(r"^\s*[-*+]\s+", lines[i]) and not re.match(r"^\s*\d+\.\s+", lines[i]):
                    item += " " + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, item)
            continue

        # numbered list
        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            flush_para()
            while i < n:
                mm = re.match(r"^\s*(\d+)\.\s+(.*)$", lines[i])
                if not mm:
                    break
                num, item = mm.group(1), mm.group(2)
                i += 1
                while i < n and lines[i].strip() and lines[i].startswith("  ") and not re.match(r"^\s*\d+\.\s+", lines[i]) and not re.match(r"^\s*[-*+]\s+", lines[i]):
                    item += " " + lines[i].strip()
                    i += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.9)
                p.paragraph_format.first_line_indent = Cm(-0.6)
                p.paragraph_format.space_after = Pt(3)
                add_inline(p, f"{num}.\u00a0 " + item)
            continue

        # ordinary paragraph line (accumulate)
        para_buf.append(line)
        i += 1

    flush_para()
    doc.save(docx_path)


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "Lighthouse-and-Plan-v2.1.md"
    dst = sys.argv[2] if len(sys.argv) > 2 else "Lighthouse-and-Plan-v2.1.docx"
    convert(src, dst)
    print("wrote", dst)
