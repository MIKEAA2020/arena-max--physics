#!/usr/bin/env python3
"""Render the fixed-seed chi128 scan without assigning a phase order.
Convergence, same-phase similarity, finite-chi branches and thermodynamic
coexistence are deliberately distinct. No coarse-grid sign test is a verdict.
"""
from pathlib import Path
import csv, json, math, sys, importlib.util
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent
R=ROOT/'v2.2/verify/replica';path=R/'n3_chi128_scan.json'
if not path.exists():raise SystemExit('No scan record yet')
db=json.loads(path.read_text());protocol=json.loads((HERE/'chi128-scan-protocol.json').read_text())
rows=[];pairs=0;new_runs=0
for p in protocol['grid']:
 case=db.get('cases',{}).get(f'{p:.7f}',{})
 a=case.get('ordered',{});b=case.get('disordered',{})
 if a.get('converged') and b.get('converged'):pairs+=1
 if not case.get('reused_verified_reference'):
  new_runs+=sum(bool(r) and r.get('status') not in ['running','pending'] for r in [a,b])
 for side,r in [('ordered',a),('disordered',b)]:
  row=dict(p=p,preparation=side,status=r.get('status','pending'),accepted=bool(r.get('converged')),reference_reused=case.get('reused_verified_reference',False))
  row.update({k:r.get(k) for k in ['eps','iters','total_stage_seconds','f_half','f_full_rayleigh','m_species','m_parity','color_l2_order','xi_chi','checkpoint','checkpoint_sha256']})
  rows.append(row)
fields=['p','preparation','status','accepted','reference_reused','eps','iters','total_stage_seconds','f_half','f_full_rayleigh','m_species','m_parity','color_l2_order','xi_chi','checkpoint','checkpoint_sha256']
with (HERE/'chi128-scan.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
def fmt(x,d=6):
 return '—' if x is None or not math.isfinite(float(x)) else f'{float(x):.{d}f}'
def sci(x):return '—' if x is None else f'{x:.2e}'
text=['# χ128 six-rate, two-preparation scan','',
'16 October 2026 · A-1 continuation, not a new Lighthouse clue.','',
f"**Campaign status:** {db.get('status')}. **Accepted paired rates:** {pairs}/6. **New terminal preparation records:** {new_runs}/10; the two previously accepted p=.30486 states are reused, not reoptimized.",'',
'## Method, allocation and scope','',
'The grid is p={.3040,.3045,.3047,.30486,.3050,.3055}. Each target uses the same two immutable χ128 warm sources as the verified single-point experiment, with the same Hermitian half-row optimizer and 1e-7 residual tolerance. Thus this is a fixed-initial-state matrix, **not a bidirectional continuation/hysteresis experiment**. The old sources are initializers, not accepted phase observations. No χ96 production or automatic χ192 is scheduled.','',
'Authorized allocation: at most two single-thread CPU workers, 90 minutes per paired new rate, up to 7.5 hours overall. Real AL/AR/AC/C arrays are atomically checkpointed; FL/FR are omitted caches, recomputed on a restart. This serialization policy was checked against the unchanged optimizer. Compact review ZIPs omit the bulk scan states; the full repository preserves them.','',
'Session recovery is recorded separately when processes do not survive a pause. Real partial states, prior JSON segments and source hashes are retained; recomputable environments are rebuilt. Checkpoint timings are lower bounds for saved optimizer segments, not an exact total including lost work or downtime. The original overall campaign deadline is not silently reset.','',
'## Results','',
'Preparation labels describe initial histories, not guaranteed physical phases. Only accepted rows enter observable comparisons. Missing values are not zero. Optimizer residuals and contraction defects are not global free-energy error bounds.','',
'| p | Preparation | Status | Residual | f_period_Rayleigh | m_species | Color L2 | xi_chi |','|---:|---|---|---:|---:|---:|---:|---:|']
for r in rows:
 text.append(f"| {r['p']:.7f} | {r['preparation']} | {r['status']} | {sci(r['eps'])} | {fmt(r['f_full_rayleigh'],10) if r['accepted'] else '—'} | {fmt(r['m_species']) if r['accepted'] else '—'} | {fmt(r['color_l2_order']) if r['accepted'] else '—'} | {fmt(r['xi_chi'],2) if r['accepted'] else '—'} |")
text+=['','Color L2 is a symmetry-invariant norm of the six one-point amplitudes about their uniform value, not a claim that all signed-coordinate amplitudes are probabilities. It avoids confusing orientation-dependent species components with different phases.','',
'## Same-period differences and symmetry alignment','','| p | f_period,O − f_period,D | Largest evaluated aligned fidelity/site | Orientation search |','|---:|---:|---:|---|']
for p in protocol['grid']:
 c=db.get('cases',{}).get(f'{p:.7f}',{});a=c.get('ordered',{});b=c.get('disordered',{});q=c.get('comparison',{})
 delta=a['f_full_rayleigh']-b['f_full_rayleigh'] if a.get('converged') and b.get('converged') else None
 text.append(f"| {p:.7f} | {sci(delta)} | {fmt(q.get('largest_evaluated_fidelity_per_site'),10)} | {q.get('status','not applicable / pending')} |")
observed=db.get('cases',{}).get('0.3047000',{})
if observed.get('ordered',{}).get('converged') and observed.get('disordered',{}).get('converged'):
    a,b=observed['ordered'],observed['disordered'];q=observed.get('comparison',{})
    text+=['',f"At p=.3047 the accepted histories have m_species {a['m_species']:.6f} / {b['m_species']:.6f} and parity {a['m_parity']:.6f} / {b['m_parity']:.6f}. Their invariant order and symmetry-aligned comparison must be used together: opposite parity signs alone do not identify different thermal phases. Largest evaluated aligned fidelity/site is {fmt(q.get('largest_evaluated_fidelity_per_site'),10)}. These large-order solutions are consistent with symmetry-related orientations, not evidence by themselves of ordered/disordered coexistence."]
text+=['','## What would count as coexistence evidence','',
'**A finite bulk free-energy gap is not coexistence.** At a genuine first-order coexistence point, the competing bulk free-energy densities agree. A non-vanishing observable/slope jump or positive interface cost is a different quantity. Distinct stationary branches can persist off coexistence; optimizers need not merge immediately on either side.','',
'Small sign changes of f_period,O−f_period,D between nearly equivalent solutions do not identify a branch crossing: O/D are initialization labels, not certified ordered/disordered branches. Symmetry-related ordered orientations also have equal bulk pressure away from a thermal transition.','',
'This coarse six-point scan can locate a change, identify finite-χ competing branches or bracket a candidate crossing. It cannot guarantee that a narrow metastability window was sampled. A candidate requires distinct symmetry-inequivalent stationary branches, a properly tuned same-functional crossing, and χ/size-systematic control of any non-vanishing discontinuity. Multiple finite-χ minima alone do not establish thermodynamic first order; no separation on this grid does not prove continuity.','',
'Numerical stationarity is for the half-row objective. Full-period values are evaluations on those states, not independently optimized full-period branches. Half/full finite-χ differences, truncation and optimizer systematics remain explicit.','',
'The previous latent-heat proxy near .11 and physical “coexistence length 40–60” remain withdrawn/unvalidated. Reported xi_chi is an MPS length, not a certified physical coexistence length. The χ128 single-point result did not isolate χ from runtime and initialization, and does not prove that every earlier large-order χ96 trial was merely an artifact or that p=.30486 is on the thermodynamic disordered side.','',
'**Transition-order verdict: unresolved.** No automatic classifier promotes a finite gap, a small residual, or a lack of sampled branch separation to a verdict.','',
'## Reproduction, framework and publication','',
'See `chi128-scan-protocol.json`, `SCAN128-ASSESSMENT.md`, raw `n3_chi128_scan.json`, per-case JSON/logs, and the actual NPZ states. Qur’anic/Lighthouse governance retains real, dependent creation distinct from the Creator; this interpretation changes neither the replica operator nor an empirical result. It is not inferred from phase order. The v34 review continuation makes that scope explicit; original v33 remains preserved.','',
'Publication uses ordinary non-force commits, with exact-secret/nested-archive scanning. The PAT is private workspace data, excluded from every repository snapshot and review archive. Source commit receipts and final remote verification are separate from numerical convergence.','']
(HERE/'SCAN128-RESULTS.md').write_text('\n'.join(text))
# Standalone exact-data SVG, no fabricated fit or error bars.
W,H=960,480;L,T,RR,B=78,55,920,390;x0,x1=.3039,.3056
vals=[r['color_l2_order'] for r in rows if r['accepted'] and r['color_l2_order'] is not None];ymax=max(1.,max(vals,default=0)*1.1)
sx=lambda p:L+(p-x0)/(x1-x0)*(RR-L);sy=lambda y:B-y/ymax*(B-T)
svg=[f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">','<rect width="100%" height="100%" fill="white"/>','<g font-family="Arial,sans-serif" fill="#173047">','<text x="28" y="28" font-size="20" font-weight="bold">χ128 fixed-source scan: accepted one-point order diagnostics</text>']
for y in [0,.25,.5,.75,1.]:
 yy=sy(y);svg+=[f'<line x1="{L}" y1="{yy}" x2="{RR}" y2="{yy}" stroke="#e4eaf0"/>',f'<text x="65" y="{yy+4}" text-anchor="end" font-size="12">{y:g}</text>']
for p in protocol['grid']:svg.append(f'<text x="{sx(p)}" y="418" text-anchor="middle" font-size="11">{p:.5f}</text>')
for side,color,offset in [('ordered','#126b9a',-2),('disordered','#bd5b2e',2)]:
 for r in rows:
  if r['preparation']==side and r['accepted'] and r['color_l2_order'] is not None:
   svg.append(f'<circle cx="{sx(r["p"])+offset}" cy="{sy(r["color_l2_order"])}" r="5" fill="{color}"/>')
 yy=76 if side=='ordered' else 101;svg+=[f'<circle cx="100" cy="{yy}" r="5" fill="{color}"/>',f'<text x="115" y="{yy+4}" font-size="13">{side}-start history</text>']
svg+=['<text x="490" y="447" text-anchor="middle" font-size="14">Measurement rate p</text>','<text x="24" y="240" transform="rotate(-90 24 240)" text-anchor="middle" font-size="13">Invariant color L2 order</text>','<text x="78" y="473" font-size="12">Accepted points only; no fit, global error bars, or transition-order certificate.</text>','</g></svg>']
(HERE/'chi128-scan.svg').write_text('\n'.join(svg))
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
module.convert(str(HERE/'SCAN128-RESULTS.md'),str(HERE/'chi128-scan-report.docx'))
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
output=HERE/'chi128-scan-report.docx';doc=Document(output)
for section in doc.sections:
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.left_margin=Cm(1.8);section.right_margin=Cm(1.8)
    section.header.paragraphs[0].text='A-1 / χ128 FIXED-SOURCE RATE SCAN'
    for run in section.header.paragraphs[0].runs:run.font.size=Pt(8)
    footer=section.footer.paragraphs[0];footer.text='Author-review research / finite-χ observations / '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for run in footer.runs:run.font.size=Pt(8)
for table in doc.tables:
    for i,row in enumerate(table.rows):
        props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
        if i==0:props.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            if i==0:module.set_cell_shading(cell,'E6EEF5')
            elif i%2==0:module.set_cell_shading(cell,'F5F7FA')
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(2)
                for run in paragraph.runs:run.font.size=Pt(8 if len(table.columns)>=8 else 9)
doc.core_properties.title='χ128 fixed-source six-rate scan'
doc.core_properties.author='Arena research assistant';doc.save(output)
print(json.dumps(dict(status=db.get('status'),accepted_paired_rates=pairs,new_terminal_records=new_runs)))
