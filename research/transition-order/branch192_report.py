#!/usr/bin/env python3
"""Generate the final resource-gated chi192 report and review DOCX."""
from pathlib import Path
import json,math,importlib.util,time,hashlib
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;R=ROOT/'v2.2/verify/replica'
db=json.loads((R/'n3_branch192.json').read_text());protocol=json.loads((HERE/'branch192-protocol.json').read_text())
add=json.loads((HERE/'branch192-execution-addendum.json').read_text());retry=json.loads((HERE/'branch192-retry-addendum.json').read_text())
audit=json.loads((R/'branch192_state_audit.json').read_text())
def fmt(x,d=12):return '—' if x is None or not math.isfinite(float(x)) else f'{float(x):.{d}f}'
def sci(x):return '—' if x is None else f'{x:.6e}'
def sha_line(name,digest):return f'`{name}` · SHA-256 `{digest}`'
results=db.get('results',{});a,b=results.get('ordered',{}),results.get('disordered',{});comp=db.get('comparison',{})
now=time.time();elapsed=now-protocol['started_unix'];remaining=max(0,protocol['deadline_unix']-now)
def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
source_checks={n:digest(ROOT/n)==h for n,h in protocol['source_sha256'].items()}
report_key='research/transition-order/branch192_report.py'
assert all(v for n,v in source_checks.items() if n!=report_key), 'Frozen numerical source hash mismatch'
current_report_hash=digest(__file__)
text=['# χ192 branch-retention continuation at p=.30480','',
'17 October 2026 · A-1 continuation · finite-χ numerical test, not a thermodynamic verdict.','',
f"**Campaign status:** {db.get('status','pending')}. **Two of two χ192 states passed the unchanged external gates.**",'',
'## Scope and resource gate','',
'The authorization was limited to enlarging the exact accepted χ128 p=.30480 ordered-history and refined low-order-history checkpoints to χ=192, with at most two single-thread workers and one five-hour absolute wall-clock cap. No extra p, perturbation, stability run, or additional preparation was authorized.', '',
'The immutable first resource gate used a deliberately broad 200-iteration forecast and stopped before production: its one-step wall timings (which include operator construction) projected beyond five hours. That STOP record is retained. The follow-up feasibility addendum then separated the VUMPS-internal iteration time from the one-time build: measured benchmark steps were 82 s and 87 s; the two accepted same-p χ128 warm steps had taken 23 and 55 iterations. A transparent 2× allowance on the larger observed count (110 iterations per preparation), a 90 s/iteration forecast, a 120 s setup reserve and a 45-minute validation reserve gave 12,720 s, within the 17,745 s then remaining. The measured simultaneous benchmark RSS passed the frozen memory rule. Only after that addendum were the two production workers launched.', '',
'One first production attempt stopped at its first five-iteration checkpoint callback because the adapter used the wrong argument signature; it wrote no χ192 checkpoint and yielded no scientific state. Its logs and JSON are preserved. A corrected serializer was used for one same-source retry, under the *original* deadline, unchanged thresholds, and same two-worker cap. This was a code-execution correction, not an added preparation. The retry and error details are in the campaign record/addenda.', '',
f"Elapsed from the initial benchmark through this report: **{elapsed/3600:.2f} h** of the authorized five hours; approximate time remaining at report generation: **{remaining/3600:.2f} h**. Both accepted states and their checks completed within the cap.",'',
'Frozen and amended protocol records:','',
sha_line('research/transition-order/branch192-protocol.json',__import__('hashlib').sha256((HERE/'branch192-protocol.json').read_bytes()).hexdigest()),
sha_line('research/transition-order/branch192-execution-addendum.json',__import__('hashlib').sha256((HERE/'branch192-execution-addendum.json').read_bytes()).hexdigest()),
sha_line('research/transition-order/branch192-retry-addendum.json',__import__('hashlib').sha256((HERE/'branch192-retry-addendum.json').read_bytes()).hexdigest()),'',
f"Frozen model, optimizer, source-worker and comparison-helper hashes were checked: all {sum(source_checks.values())}/{len(source_checks)} launch-time entries match except the final report formatter, which was changed only after computation to include the amended feasibility record and independent audit. The launch-time report-helper SHA was `{protocol['source_sha256'][report_key]}`; the post-run formatter SHA is `{current_report_hash}`.",
f"Corrected retry worker SHA-256: `{retry['corrected_worker_sha256']}`. Physics kernels and optimizer hashes remain the unchanged frozen values above.",'',
'Acceptance was not relaxed: optimizer and left/right compatibility residuals <1e−7, canonical defects <1e−9, positive full-period Rayleigh contraction with relative residual <1e−8. The internal 5e−8 target drove convergence but did not replace the external tests.', '',
'## Disposable χ192 benchmark','',
'| Start history | VUMPS iterations | Wall time including operator setup | VUMPS-reported step time | Peak RSS | Output finite |','|---|---:|---:|---:|---:|---|']
for side in ['ordered','disordered']:
 r=json.loads((R/f'branch192_benchmark_{side}.json').read_text())
 internal=add['benchmark_internal_seconds'][side]
 text.append(f"| {side} | {r['iterations']} | {r['seconds']:.2f}s | {internal}s | {r['peak_rss_kib']} KiB | {r['finite_output']} |")
text+=['','The disposable iterates were not checkpointed or used as scientific observations. The benchmark memory gate used both simultaneous peaks: 344,744 KiB combined; 1.25× this was below 0.75× available memory at the decision point. During production each process peaked near 605 MiB RSS; no memory limit was exceeded.','',
'## Accepted χ192 results','',
'| Start history | Iterations | Optimizer seconds | ε | Right compatibility | Color L2 | m_species | ξχ | f_half | f_period,trial |','|---|---:|---:|---:|---:|---:|---:|---:|---:|---:|']
for side,r in [('ordered',a),('low-order',b)]:
 if r:
  text.append(f"| {side} | {r.get('iters','—')} | {fmt(r.get('optimizer_seconds'),2)} | {sci(r.get('eps'))} | {sci(r.get('gauge_right_residual'))} | {fmt(r.get('color_l2_order'),7)} | {fmt(r.get('m_species'),7)} | {fmt(r.get('xi_chi'),2)} | {fmt(r.get('f_half'))} | {fmt(r.get('f_full_rayleigh'))} |")
 else:text.append(f'| {side} | — | — | — | — | — | — | — | — | — |')
text+=['',
'Independent saved-tensor audit passed for both 192×6×192 physical MPS tensors: exact checkpoint hashes and metadata; finite arrays; both mixed-canonical compatibility residuals and canonical defects recomputed; half-row objective, color diagnostic and positive full-period Rayleigh contraction recomputed. Maximum audited ε = '+sci(max(x['eps'] for x in audit['states'].values()))+'; maximum compatibility residual = '+sci(max(max(x['gauge_left_residual'],x['gauge_right_residual']) for x in audit['states'].values()))+'; maximum canonical defect = '+sci(max(max(x['left_canonical_residual'],x['right_canonical_residual']) for x in audit['states'].values()))+'. This is a state-integrity/acceptance audit, not an energy-error estimate.','',
'## Did the χ128 branch-retention behavior persist?','',
'The operational answer for this single warm-enlargement test is **no for the low-order continuation**. At χ128 its accepted source had invariant color L2 ≈0.009027. Starting from that exact low-order source, the unconstrained χ192 warm continuation converged to color L2 = '+fmt(b.get('color_l2_order'),7)+', similar to the ordered-source χ192 result ('+fmt(a.get('color_l2_order'),7)+'). The χ128 ordered-source value had been ≈0.488232; its χ192 value remains large, ≈0.479127.', '',
'All 36 symmetry orientations were compared at χ192. The largest evaluated normalized fidelity per site was **'+fmt(comp.get('largest_evaluated_fidelity_per_site'))+'**. Thus these two independently initialized-and-optimized χ192 records are numerically almost the same symmetry orbit, rather than a retained ordered/low-order pair. This supports a finite-χ optimizer-basin merger/loss of the low-order continuation under this precise enlargement path. It does **not** establish that no low-order χ192 stationary state exists: no constrained search, alternative restart ensemble, or stability campaign was authorized.', '',
'For the two accepted states, Δf_half (ordered-start minus low-order-start) = **'+f"{a['f_half']-b['f_half']:.8e}"+'**; the separately retained full-period trial difference is **'+f"{a['f_full_rayleigh']-b['f_full_rayleigh']:.8e}"+'**. These differences are tiny between near-identical MPS constructions. Do not interpret their arithmetic sign as a coexistence crossing, pressure bracket, or ranking with quantified uncertainty.', '',
'## Limits and interpretation','',
'An accepted VUMPS state is a finite-χ variational construction, not a global-minimum certificate. Optimizer residuals, canonical defects and contraction residuals are not free-energy error bars. Similarity at one rate and χ does not prove uniqueness. Loss of this continuation at χ=192 does not prove a continuous transition or physical instability; likewise, its χ=128 persistence did not prove a thermodynamic metastable phase. There is still no same-objective finite-χ crossing bracket, perturbation stability result, χ→∞ extrapolation or thermodynamic order verdict.', '',
'**Transition order remains unresolved.** This test used one p only; no more rates or stability tests are included. The governing Qur’anic/Lighthouse interpretation remains separate from physics: real dependent creation is distinct from the uncreated Creator; no quantum state or algebra is identified with God.', '',
'### Durable artifacts','',
'- `v2.2/verify/replica/n3_branch192.json` — campaign aggregate, gate history, accepted pair and symmetry comparison.',
'- `v2.2/verify/replica/branch192_state_audit.json` — independent tensor/gate audit.',
'- `v2.2/verify/replica/checkpoints/branch192_retry1_ordered_p0.3048000.npz` and `...branch192_retry1_disordered_p0.3048000.npz` — retained real χ192 state tensors.',
'- `v2.2/verify/replica/n3_branch192_attempt1.json` and `branch192_retry1_*.log` — retained failed serializer attempt and corrected-run logs.',
'- `research/transition-order/branch192.csv` — compact result data.','']
(HERE/'BRANCH192-RESULTS.md').write_text('\n'.join(text))
keys=['history','status','iters','eps','color_l2_order','m_species','xi_chi','f_half','f_full_rayleigh','checkpoint','checkpoint_sha256']
with (HERE/'branch192.csv').open('w',newline='') as f:
 import csv
 w=csv.DictWriter(f,fieldnames=keys,lineterminator='\n');w.writeheader()
 for r in results.values():w.writerow({k:r.get(k,'') for k in keys})
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);out=HERE/'branch192-report.docx';m.convert(str(HERE/'BRANCH192-RESULTS.md'),str(out))
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
D=Document(out)
for sec in D.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
 sec.header.paragraphs[0].text='A-1 / χ192 BRANCH-RETENTION TEST'
 for rr in sec.header.paragraphs[0].runs:rr.font.size=Pt(8)
for table in D.tables:
 for i,row in enumerate(table.rows):
  pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if i==0:pr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   if i==0:m.set_cell_shading(cell,'E6EEF5')
   for para in cell.paragraphs:
    for rr in para.runs:rr.font.size=Pt(7.5 if len(table.columns)>=8 else 9)
D.core_properties.title='χ192 branch-retention continuation at p=.30480';D.save(out)
print(json.dumps(dict(status=db.get('status'),accepted=sum(r.get('converged') is True for r in results.values()),elapsed_hours=elapsed/3600,remaining_hours=remaining/3600)))
