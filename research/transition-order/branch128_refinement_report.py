#!/usr/bin/env python3
"""Report the separately authorized refinement without rewriting its failed parent."""
from pathlib import Path
import json,math,importlib.util
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;R=ROOT/'v2.2/verify/replica'
r=json.loads((R/'n3_branch128_refinement.json').read_text());pilot=json.loads((R/'n3_branch128_pilot.json').read_text())
o=pilot['stages']['2']['ordered'];old=pilot['stages']['2']['disordered']
def fmt(x,d=12):return '—' if x is None or not math.isfinite(float(x)) else f'{float(x):.{d}f}'
def sci(x):return '—' if x is None else f'{x:.9e}'
text=['# χ128 same-point convergence-safety refinement','',
'16 October 2026 · A-1 · a separately recorded numerical amendment, not a new rate.','',
f"**Refinement status:** {r['status']}. **Original pilot status remains:** {pilot['status']}. The original failed record and checkpoint are not overwritten or retroactively accepted.",'',
'## Why refinement was authorized','',
f"At p=.30480, the original small-order history passed the primary stopping test ({sci(old['eps'])}) but its right compatibility defect was {sci(old['gauge_right_residual'])}, above the strict 1e-7 ceiling. It was rejected, not rounded into acceptance.",'',
'The user separately authorized up to 30 minutes at the same p=.30480, χ128 and unbiased operator. Only the internal convergence-driving goal is tightened to 5e-8. External acceptance is unchanged: primary and BOTH compatibility residuals <1e-7, canonical defects <1e-9, positive full-period contraction and relative defect <1e-8. No phase constraint, new p, perturbation stability production, crossing bracket or χ192 is added.','',
'An internal goal is not a new external threshold. If an internal driving budget ends, the saved state can be accepted only after all original external gates and period validation are actually checked within the reserved time. Its internal-target and budget status remain explicit. A changed order character is an outcome, not prevented.','',
'## Refinement geometry and acceptance','','| Diagnostic | Refined small-order history |','|---|---:|',
f"| Primary optimizer residual | {sci(r.get('eps'))} |",
f"| Left compatibility | {sci(r.get('gauge_left_residual'))} |",
f"| Right compatibility | {sci(r.get('gauge_right_residual'))} |",
f"| Left canonical defect | {sci(r.get('left_canonical_residual'))} |",
f"| Right canonical defect | {sci(r.get('right_canonical_residual'))} |",
f"| Internal 5e-8 goal met | {r.get('internal_goal_met','pending')} |",
f"| Period contraction defect | {sci(r.get('period_contraction_relative_residual'))} |",
f"| Total refinement seconds | {fmt(r.get('total_refinement_seconds'),2)} |"]
if r.get('converged'):
 c=r.get('common_comparison',{})
 text+=['','## Augmented comparison at the same p=.30480','',
 'The large-order result is the already accepted original-pilot state. The other column is the separately validated refinement. These are finite-χ numerical states; history labels are not automatically thermodynamic phases.','',
 '| Diagnostic | Accepted large-order history | Refined small-order history |','|---|---:|---:|',
 f"| f_half (stationary objective estimate) | {fmt(o['f_half'])} | {fmt(r['f_half'])} |",
 f"| f_period_Rayleigh (separate trial evaluation) | {fmt(o['f_full_rayleigh'])} | {fmt(r['f_full_rayleigh'])} |",
 f"| m_species | {fmt(o['m_species'],8)} | {fmt(r.get('m_species'),8)} |",
 f"| abs(m_parity) | {fmt(abs(o['m_parity']),8)} | {fmt(abs(r['m_parity']),8)} |",
 f"| Invariant color L2 | {fmt(o['color_l2_order'],8)} | {fmt(r.get('color_l2_order'),8)} |",
 f"| xi_chi | {fmt(o['xi_chi'],4)} | {fmt(r.get('xi_chi'),4)} |",
 '',f"Ordered-history minus refined-history: **Δf_half={o['f_half']-r['f_half']:.8e}** and **Δf_period,trial={o['f_full_rayleigh']-r['f_full_rayleigh']:.8e}**. These are different finite-χ functionals, not one mixed crossing or global error bounds.",
 '',f"Symmetry-aligned comparison: {c.get('status','not performed')}; largest evaluated fidelity/site {fmt(c.get('largest_evaluated_fidelity_per_site'))}. This is a constructed-state diagnostic, not a phase-order classifier."]
else:text+=['','No fully accepted augmented common-target pair is recorded. Rejected/partial fields are not promoted to stationary branch values.']
text+=['','## What is still missing','',
'A nonzero energy difference at one common p is not coexistence. Competing bulk branch free energies must agree at a properly bracketed crossing of distinct phases. This amendment provides neither a second common rate nor an independently tested local/thermodynamic stability certificate. The half-row objective and full-period trial evaluation differ at finite χ; no global variational/truncation error bound or χ→infinity control is obtained here.','',
'**Physical transition order remains unresolved.** The earlier pressure values were already measured; the missing work concerns identification/retention, crossing, stability and systematic errors of phase branches. No numerical result is forced by the governing creator–creation interpretation.','']
(HERE/'BRANCH128-REFINEMENT-RESULTS.md').write_text('\n'.join(text))
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);out=HERE/'branch128-refinement-report.docx';m.convert(str(HERE/'BRANCH128-REFINEMENT-RESULTS.md'),str(out))
from docx import Document
from docx.shared import Cm,Pt
D=Document(out)
for s in D.sections:s.page_width=Cm(21);s.page_height=Cm(29.7);s.left_margin=Cm(1.8);s.right_margin=Cm(1.8)
for table in D.tables:
 for i,row in enumerate(table.rows):
  for cell in row.cells:
   if i==0:m.set_cell_shading(cell,'E6EEF5')
   for p in cell.paragraphs:
    for x in p.runs:x.font.size=Pt(9)
D.core_properties.title='χ128 same-point convergence-safety amendment';D.save(out)
print(json.dumps(dict(status=r['status'],accepted=r.get('converged'),original_pilot_status=pilot['status'])))
