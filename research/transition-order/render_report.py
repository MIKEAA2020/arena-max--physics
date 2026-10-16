#!/usr/bin/env python3
"""Render the continuation report as modern OOXML with real source hyperlinks."""
from pathlib import Path
import importlib.util,re,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
base=m.add_inline;rx=re.compile(r'\[([^\]]+)\]\(([^)]+)\)')
def rich(p,text,base_bold=False,base_italic=False,size=None,font=None):
    pos=0
    for hit in rx.finditer(text):
        base(p,text[pos:hit.start()],base_bold,base_italic,size,font)
        label,url=hit.groups()
        if not url.startswith(('http://','https://')):url='https://github.com/MIKEAA2020/arena-max--physics/blob/main/research/transition-order/'+url
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'176A9E');pr.append(color)
        u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
        if size:
            z=OxmlElement('w:sz');z.set(qn('w:val'),str(int(size*2)));pr.append(z)
        r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);h.append(r);p._p.append(h);pos=hit.end()
    base(p,text[pos:],base_bold,base_italic,size,font)
m.add_inline=rich
out=HERE/'transition-order-report.docx';m.convert(str(HERE/'RESULTS.md'),str(out))
d=Document(out)
for sec in d.sections:
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
    sec.header.paragraphs[0].text='A-1  /  COMPUTATIONAL CONTINUATION'
    for r in sec.header.paragraphs[0].runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('627185')
    foot=sec.footer.paragraphs[0];foot.text='16 October 2026  •  Author-review research  •  '
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for r in foot.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor.from_string('627185')
for table in d.tables:
    for i,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr();nosplit=OxmlElement('w:cantSplit');trpr.append(nosplit)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for cell in row.cells:
            if i==0:m.set_cell_shading(cell,'E6EEF5')
            elif i%2==0:m.set_cell_shading(cell,'F5F7FA')
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2)
                for r in p.runs:r.font.size=Pt(8 if len(table.columns)>=8 else 8.5)
d.core_properties.title='Transition-order continuation: executed work and status'
d.core_properties.subject='Exact replica operator, transfer MPS, feasibility gates and verification'
d.core_properties.author='Arena research assistant'
d.save(out)
# Native OOXML structure/content round trip; no claim of a submission-layout PDF.
check=Document(out);assert len(check.tables)>=5 and 'Decision and outcome' in '\n'.join(p.text for p in check.paragraphs)
print(out)
