#!/usr/bin/env python3
"""Render branch-retention pilot data; no phase or crossing classifier."""
from pathlib import Path
import json,csv,math,importlib.util
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;R=ROOT/'v2.2/verify/replica'
db=json.loads((R/'n3_branch128_pilot.json').read_text())
rows=[]
for key,c in sorted(db.get('stages',{}).items()):
 for side in ['ordered','disordered']:
  r=c.get(side,{})
  if r:rows.append(dict(stage=key,history=side,**r))
fields=['stage','history','p','status','converged','eps','iters','total_stage_seconds','f_half','f_full_rayleigh','m_species','m_parity','color_l2_order','xi_chi','checkpoint','checkpoint_sha256']
with (HERE/'branch128-pilot.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader();w.writerows(rows)
def fmt(x,d=8):return '—' if x is None or not math.isfinite(float(x)) else f'{float(x):.{d}f}'
def sci(x):return '—' if x is None else f'{x:.3e}'
text=['# χ128 unbiased branch-following pilot','',
'16 October 2026 · A-1 continuation · numerical retention test, not an order verdict.','',
f"**Status:** {db.get('status')}. The user separately authorized at most three hours overall, with at most two single-thread CPU workers. No χ192, extra crossing point or perturbation production is automatic.",'',
'## Protocol and objective','',
'Large-order history: accepted p=.30470 → .30475 → .30480. Small-order history: accepted p=.30486 → .30483 → .30480. The stage-1 targets differ, so their energies must not be compared as a same-p branch gap. Stage 2 starts only after both stage-1 results pass the unchanged numerical gates. Every step loads the actual previous accepted MPS, not a cached scalar or a reset to the old grid initializer.','',
'The objective is the unchanged Hermitian half-row transfer-MPS functional, with residual and left/right compatibility below 1e-7, canonical checks below 1e-9 and full-period contraction defect below 1e-8. No phase, order-parameter or pinning constraint is imposed. A history can lose its original character; it is not forced to survive as a branch.','',
'Half-row stationarity and full-period trial evaluation are different finite-χ objects. Both functionals are reported separately; do not mix them to infer a crossing. A local stopping test does not certify a global minimum, a physical spinodal or stability against all perturbations.','',
'## Accepted anchors','',
'| History | Anchor p | Color L2 | f_half | f_period_Rayleigh |','|---|---:|---:|---:|---:|']
for side,r in db['anchors'].items():text.append(f"| {side}-start history | {r['p']:.5f} | {fmt(r['color_l2_order'],6)} | {fmt(r['f_half'],12)} | {fmt(r['f_full_rayleigh'],12)} |")
text+=['','These anchors are at different rates and are not a coexistence pair. Their original accepted checkpoint hashes are retained.','',
'## Continuation steps','','| Stage | History | p | Status | Residual | Color L2 | f_half | f_period_Rayleigh |','|---:|---|---:|---|---:|---:|---:|---:|']
for r in rows:
 good=r.get('converged') is True
 text.append(f"| {r['stage']} | {r['history']} | {r['p']:.5f} | {r['status']} | {sci(r.get('eps'))} | {fmt(r.get('color_l2_order'),6) if good else '—'} | {fmt(r.get('f_half'),11) if good else '—'} | {fmt(r.get('f_full_rayleigh'),11) if good else '—'} |")
if not rows:text.append('| — | — | — | pending | — | — | — | — |')
text+=['','Missing or rejected results are not zero and are not stationary branch observations. If a step fails, its path stops; it is not promoted by skipping to the next target.','',
'## Common-target comparison','']
c=db.get('stages',{}).get('2',{})
if c.get('ordered',{}).get('converged') and c.get('disordered',{}).get('converged'):
 a,b=c['ordered'],c['disordered'];q=db.get('common_comparison',{})
 assert a['p']==b['p']==.3048
 text += [f"At p=.30480: Δf_half (ordered-history minus disordered-history) = **{a['f_half']-b['f_half']:.6e}**; Δf_period,trial = **{a['f_full_rayleigh']-b['f_full_rayleigh']:.6e}**. These are different finite-χ functional comparisons, not error bars or a root bracket.",'',
 f"Invariant order is {a['color_l2_order']:.8f} / {b['color_l2_order']:.8f}; m_species is {a['m_species']:.8f} / {b['m_species']:.8f}; xi_chi is {a['xi_chi']:.3f} / {b['xi_chi']:.3f}.",'',
 f"Symmetry/gauge-aware comparison: {q.get('status','pending')}; largest evaluated fidelity per site {fmt(q.get('largest_evaluated_fidelity_per_site'),12)}. Similarity is not certified exact identity. Opposite parity or symmetry orientation alone is not a distinct thermal phase."]
else:text.append('No two accepted common-target results yet; no energy-gap or coexistence inference is made.')
rejected=db.get('stages',{}).get('2',{}).get('disordered',{})
if rejected and not rejected.get('converged') and rejected.get('optimizer_tolerance_passed'):
    right=rejected.get('gauge_right_residual')
    text+=['','## Exact validation failure — not rounded into acceptance','',f"The common-target small-order history passed the primary optimizer tolerance ({rejected['eps']:.12e}), but its right compatibility defect was **{right:.12e}**, above **1e-7** by {(right/1e-7-1)*100:.6f}%. The right gate was not relaxed. No full-period value was evaluated for this rejected step, and no accepted common-target branch comparison was constructed. This is a numerical gate failure, not evidence of a physical phase or spinodal.",
    '', 'A same-point safety refinement with a stricter internal goal and unchanged external gates is described in CONVERGENCE-REFINEMENT.md. It requires separate authorization; the original failed attempt remains preserved.']
text+=['','## Limits and next gates','',
'This pilot samples continuity/retention from accepted anchors. It does not compute exact continuous free-energy functions. If both histories reach a similar family, that does not prove continuity; if distinct finite-χ states survive, that does not prove first order. A nonzero trial free-energy gap at one p is not coexistence.','',
'A candidate crossing requires retained symmetry-inequivalent stationary families at at least another common rate, a controlled same-objective pressure bracket, and stability/χ/size-systematic analysis of a non-vanishing discontinuity. Perturb-and-return checks are not performed by this allocation; an exact fixed point remaining fixed is not a stability test. No physical spinodal or latent heat is certified.','',
'For an ordinary continuous transition f′ is continuous while higher derivatives can be singular. A kink in f normally means a first-derivative jump in a generic tuning direction. Sharpening with χ can occur for continuous transitions as finite-entanglement rounding recedes.','',
'**Thermodynamic order remains unresolved.** The original six-rate data, v33 and the v34 creator–creation interpretive governance remain preserved. No numerical result is inferred from theology or vice versa.','']
(HERE/'BRANCH128-PILOT-RESULTS.md').write_text('\n'.join(text))
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
output=HERE/'branch128-pilot-report.docx';m.convert(str(HERE/'BRANCH128-PILOT-RESULTS.md'),str(output))
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
D=Document(output)
for sec in D.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
 sec.header.paragraphs[0].text='A-1 / χ128 UNBIASED CONTINUATION PILOT'
 for r in sec.header.paragraphs[0].runs:r.font.size=Pt(8)
for table in D.tables:
 for i,row in enumerate(table.rows):
  props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
  if i==0:props.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   if i==0:m.set_cell_shading(cell,'E6EEF5')
   for p in cell.paragraphs:
    for r in p.runs:r.font.size=Pt(8 if len(table.columns)>=8 else 9)
D.core_properties.title='χ128 unbiased branch-following pilot';D.save(output)
print(json.dumps(dict(status=db['status'],recorded_steps=len(rows),accepted_steps=sum(r.get('converged') is True for r in rows))))
