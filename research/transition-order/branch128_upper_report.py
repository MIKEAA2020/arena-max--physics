#!/usr/bin/env python3
"""Report exactly one additional common rate; pressure signs are not a classifier."""
from pathlib import Path
import json,csv,math,importlib.util
ROOT=Path(__file__).resolve().parents[2];HERE=Path(__file__).resolve().parent;R=ROOT/'v2.2/verify/replica'
db=json.loads((R/'n3_branch128_upper.json').read_text());protocol=json.loads((HERE/'branch128-upper-protocol.json').read_text())
def fmt(x,d=12):return '—' if x is None or not math.isfinite(float(x)) else f'{float(x):.{d}f}'
def sci(x):return '—' if x is None else f'{x:.6e}'
text=['# χ128 second-common-rate test: p=.30485','',
'16 October 2026 · A-1 continuation · one paired rate only.','',
f"**Status:** {db['status']}. The user authorized exactly one additional common rate, warm-continued from the accepted .30480 pair. No χ192, additional rate or stability production is included.",'',
'## Prospective protocol','',
'Both accepted .30480 physical MPS states initialize the same zero-bias operator at .30485. The external gates are unchanged: primary and both compatibility residuals <1e-7, canonical defects <1e-9, positive period contraction with relative defect <1e-8. The internal driving goal is 5e-8 for a safety margin, not a new external criterion. Maximum overall allocation is 60 minutes, with two single-thread CPU workers; 55 minutes are available to optimization and time is reserved for validation/comparison.','',
'A failed gate is not rounded into a pass. If the internal driving budget ends, the actual saved state is evaluated under the original external gates; a numerical state is accepted only if those complete checks pass. No pinning field or phase constraint forces a family to persist.','',
'## Upper-rate numerical outcomes','','| History | Status | Primary residual | Right compatibility | Color L2 | m_species | xi_chi |','|---|---|---:|---:|---:|---:|---:|']
for side in ['ordered','disordered']:
 r=db.get('results',{}).get(side,{})
 good=r.get('converged') is True
 text.append(f"| {side}-start history | {r.get('status','pending')} | {sci(r.get('eps'))} | {sci(r.get('gauge_right_residual'))} | {fmt(r.get('color_l2_order'),8) if good else '—'} | {fmt(r.get('m_species'),8) if good else '—'} | {fmt(r.get('xi_chi'),4) if good else '—'} |")
text+=['','History labels are not automatically phase identities. An observed change/merger or numerical failure is not a measured physical spinodal or a bound on a thermodynamic metastability-window width.','',
'## Same-objective pressure estimates','','| p | Half-row O | Half-row D | Δf_half (O−D) | Period trial O | Period trial D | Δf_period,trial |','|---:|---:|---:|---:|---:|---:|---:|']
for p,pair in [(.3048,db['previous_common']),(.30485,db.get('results',{}))]:
 a,b=pair.get('ordered',{}),pair.get('disordered',{})
 if a.get('converged') and b.get('converged'):
  text.append(f"| {p:.5f} | {fmt(a['f_half'])} | {fmt(b['f_half'])} | {sci(a['f_half']-b['f_half'])} | {fmt(a['f_full_rayleigh'])} | {fmt(b['f_full_rayleigh'])} | {sci(a['f_full_rayleigh']-b['f_full_rayleigh'])} |")
 else:text.append(f'| {p:.5f} | — | — | no accepted pair | — | — | — |')
text+=['','The primary stationary objective is the Hermitian half-row functional. Full-period values are separate trial evaluations, not independently converged full-period branch optima. Do not mix the two columns into one pressure bracket. Last digits are numerical contraction outputs, not global error bars.','']
c=db.get('upper_comparison',{})
if c:
 text += ['## Numerical similarity and pressure-sign diagnostics','',
 f"Upper-rate symmetry comparison: {c.get('status')}; largest evaluated aligned fidelity/site **{fmt(c.get('largest_evaluated_fidelity_per_site'))}**.",'',
 f"Half-row ranking sign change between the two endpoints: **{db.get('half_objective_sign_change')}**. Period-trial ranking sign change: **{db.get('period_trial_sign_change')}**. These are arithmetic diagnostics only, not coexistence certificates."]
if all(db.get('results',{}).get(s,{}).get('converged') for s in ['ordered','disordered']):
    a,b=db['results']['ordered'],db['results']['disordered'];lo=db['previous_common']['ordered']
    if lo['color_l2_order']>.1 and max(a['color_l2_order'],b['color_l2_order'])<.01:
        text+=['','## Completed outcome: no retained two-family bracket','',
        f"Both upper-rate numerical states passed every original gate. However, the large-order history changed from invariant order {lo['color_l2_order']:.8f} at .30480 to **{a['color_l2_order']:.8f}** at .30485; the other upper-rate history has **{b['color_l2_order']:.8f}**. The upper endpoint therefore does not supply a retained large-order/small-order pair. The high aligned fidelity is a similarity diagnostic, not certified exact identity.",
        '', 'Although both tabulated ranking differences change sign, the upper O-history value now belongs to a small-order approximant. It cannot be relabeled as the ordered-phase branch pressure. Consequently **no retained large-/small-order pressure-crossing bracket was obtained** by this single-step test.',
        '', 'This does not invalidate the distinct constructed states previously retained at .30480. It says only that the selected unbiased warm step did not retain the large-order character at .30485. The test does not prove that such a branch is absent there, identify a physical spinodal, or give a thermodynamic retention-window width. A finer continuation step or different validated branch-following control would require separate authorization; none was run.']
        audit_path=HERE/'branch128-upper-data-audit.json'
        if audit_path.exists():
            audit=json.loads(audit_path.read_text());q=audit.get('actual_AL_state_diagnostics',{})
            if len(q)==2:text+=['',f"Read-only actual-AL-density audit gives upper invariant order {q['ordered']['actual_AL_color_l2']:.10f} / {q['disordered']['actual_AL_color_l2']:.10f}, confirming the small-order endpoint diagnostics independently of the center-tensor estimate. This is not a perturbation or stability test."]
text+=['','## Interpretation safeguards','',
'A candidate finite-χ pressure bracket requires two distinguishable, symmetry-inequivalent numerical families retained at BOTH endpoints and compared using the SAME objective. A sign change between nearly equivalent outputs after a history has merged is not such a branch crossing. A direct warm step that loses character does not locate a physical spinodal or prove how wide the retention window is.','',
'Even if endpoint data support a candidate finite-χ bracket, this allocation does not supply local perturbation stability, global optimality, a physical coexistence point, a thermodynamic order-parameter jump, or χ/size-extrapolated transition order. Those are separately gated. If ordering does not flip, only the ranking of these computed trial states at these rates is known; no theorem places the thermodynamic transition further up.','',
'Optimizer residuals are NOT free-energy-error bars. Comparing a pressure difference numerically to a residual does not yield a calibrated signal-to-error ratio. A lower trial functional among two states does not certify a global minimum, even at this χ. Finite-χ bias and optimizer systematics remain unbounded globally.','',
'**Thermodynamic verdict remains unresolved.** The earlier six-rate data and original stopped pilot are preserved; the accepted same-point refinement is a separate numerical amendment. The Qur’anic/Lighthouse creator–creation interpretation and mathematical statements are not changed by these numerical outcomes.','']
(HERE/'BRANCH128-UPPER-RESULTS.md').write_text('\n'.join(text))
fields=['p','history','status','converged','eps','gauge_left_residual','gauge_right_residual','f_half','f_full_rayleigh','m_species','m_parity','color_l2_order','xi_chi','checkpoint','checkpoint_sha256']
with (HERE/'branch128-upper.csv').open('w',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader()
 for p,pair in [(.3048,db['previous_common']),(.30485,db.get('results',{}))]:
  for side,row in pair.items():
   if side in ['ordered','disordered']:w.writerow(dict(row,p=p,history=side))
spec=importlib.util.spec_from_file_location('md2docx',ROOT/'v2.2/md2docx.py');m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
out=HERE/'branch128-upper-report.docx';m.convert(str(HERE/'BRANCH128-UPPER-RESULTS.md'),str(out))
from docx import Document
from docx.shared import Cm,Pt
from docx.oxml import OxmlElement
D=Document(out)
for sec in D.sections:
 sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.left_margin=Cm(1.8);sec.right_margin=Cm(1.8)
 sec.header.paragraphs[0].text='A-1 / χ128 SECOND COMMON RATE'
 for r in sec.header.paragraphs[0].runs:r.font.size=Pt(8)
for table in D.tables:
 for i,row in enumerate(table.rows):
  tr=row._tr.get_or_add_trPr();tr.append(OxmlElement('w:cantSplit'))
  if i==0:tr.append(OxmlElement('w:tblHeader'))
  for cell in row.cells:
   if i==0:m.set_cell_shading(cell,'E6EEF5')
   for p in cell.paragraphs:
    for r in p.runs:r.font.size=Pt(8 if len(table.columns)>=7 else 9)
D.core_properties.title='χ128 second-common-rate test at p=.30485';D.save(out)
print(json.dumps(dict(status=db['status'],accepted_upper_preparations=sum(r.get('converged') is True for r in db.get('results',{}).values()))))
